"""Build shared/src/gear-data.json (catalogue + Kit Spec) from the PCC gear workbook.

Usage:
  python scripts/import_gear.py --workbook "<path to PCC Gear List 2026-27 Season - v7.xlsx>"
      [--out shared/src/gear-data.json]

Teams and pools live in D1 and are managed on the admin page (/admin), not here.
Needs openpyxl. See docs/spec.md §2 for the rules.
"""

import argparse
import json
import sys
from pathlib import Path

import openpyxl

EXCLUDED_CATEGORIES = {"Senior kit", "Misc", "Other safety"}




def fail(msg: str):
    sys.exit(f"import_gear: {msg}")


def text(v) -> str:
    return "" if v is None else str(v).strip()


def header_row(rows, first: str) -> int:
    for i, r in enumerate(rows):
        if r and text(r[0]) == first:
            return i
    fail(f"no header row starting {first!r}")


def read_catalogue(ws):
    rows = list(ws.iter_rows(values_only=True))
    h = header_row(rows, "Item ID")
    cols = [text(c) for c in rows[h]]
    ci, cc, cn = cols.index("Item ID"), cols.index("Category"), cols.index("Item")
    items = []
    for r in rows[h + 1 :]:
        item_id = text(r[ci])
        if not item_id:
            continue
        cat = text(r[cc])
        if cat in EXCLUDED_CATEGORIES:
            continue
        items.append({"id": item_id, "category": cat, "name": text(r[cn])})
    return items


def read_kit_spec(ws, catalogue_ids):
    rows = list(ws.iter_rows(values_only=True))
    h = header_row(rows, "Item ID")
    cols = [text(c) for c in rows[h]]
    grade_cols = {name: i for i, name in enumerate(cols) if i >= 3 and name}
    specs = {g: {} for g in grade_cols}
    for r in rows[h + 1 :]:
        item_id = text(r[0])
        if item_id not in catalogue_ids:
            continue
        for g, i in grade_cols.items():
            v = r[i] if i < len(r) else None
            if isinstance(v, (int, float)) and v > 0:
                specs[g][item_id] = int(v)
    return specs




def build(workbook: str):
    wb = openpyxl.load_workbook(workbook, data_only=True, read_only=True)
    for name in ("Items", "Kit Spec"):
        if name not in wb.sheetnames:
            fail(f"workbook has no {name!r} sheet")
    items = read_catalogue(wb["Items"])
    specs = read_kit_spec(wb["Kit Spec"], {i["id"] for i in items})
    categories = list(dict.fromkeys(i["category"] for i in items))
    return {
        "source": f"{Path(workbook).name} — Items, Kit Spec",
        "categories": categories,
        "items": items,
        # Every grade column with quantities; the admin page offers these when adding a team.
        "specs": {g: q for g, q in specs.items() if q},
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--out", default=str(Path(__file__).resolve().parent.parent / "shared/src/gear-data.json"))
    a = ap.parse_args()
    data = build(a.workbook)
    Path(a.out).write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"{len(data['items'])} items, specs: {', '.join(data['specs'])} -> {a.out}")


if __name__ == "__main__":
    main()
